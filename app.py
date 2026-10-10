from flask import Flask, request, jsonify, render_template, send_file
from flask_cors import CORS
from groq import Groq
import os
import io
from docx import Document
import openpyxl

app = Flask(__name__, template_folder='.')
CORS(app)

# Lấy key từ biến môi trường của hệ thống (Render sẽ cung cấp)
client = Groq(api_key=os.environ.get("GROQ_API_KEY"))

# ==========================================
# CÁC HÀM XỬ LÝ ĐỌC FILE WORD / EXCEL
# ==========================================
def read_word_file(file_storage):
    try:
        doc = Document(file_storage)
        text = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(text)
    except Exception as e:
        return f"[Lỗi đọc file Word: {str(e)}]"

def read_excel_file(file_storage):
    try:
        wb = openpyxl.load_workbook(file_storage, data_only=True)
        sheet = wb.active
        text_data = []
        for row in sheet.iter_rows(values_only=True):
            if any(row):
                row_str = " | ".join([str(cell) if cell is not None else "" for cell in row])
                text_data.append(row_str)
        return "\n".join(text_data)
    except Exception as e:
        return f"[Lỗi đọc file Excel: {str(e)}]"


@app.route('/')
def home():
    return render_template('webchat.html')

@app.route('/api/chat', methods=['POST'])
def chat():
    # Hỗ trợ nhận cả dữ liệu từ FormData (có kèm file) hoặc JSON thuần (chỉ có text)
    messages_history = []
    user_message = ""
    file_content_text = ""

    if request.files.get('file'):
        file = request.files.get('file')
        user_message = request.form.get('message', '')
        filename = file.filename.lower()
        
        if filename.endswith('.docx'):
            file_content_text = f"\n\n[Nội dung từ file Word '{file.filename}':\n{read_word_file(file)}]\n"
        elif filename.endswith('.xlsx') or filename.endswith('.xls'):
            file_content_text = f"\n\n[Dữ liệu từ file Excel '{file.filename}':\n{read_excel_file(file)}]\n"
        
        # Tạo lại lịch sử tin nhắn dạng đơn giản khi có gửi file
        full_prompt = user_message + file_content_text
        messages_history = [{"role": "user", "content": full_prompt}]
    else:
        # Trường hợp chat thông thường bằng JSON
        data = request.get_json(silent=True) or {}
        messages_history = data.get('messages', [])

    if not messages_history:
        return jsonify({'reply': 'Vui lòng nhập nội dung hoặc đính kèm tài liệu.'})

    # System Prompt tối ưu: Cực kỳ ngắn gọn, sắc bén, đúng trọng tâm và chuẩn hóa toán học
    system_prompt = {
        "role": "system", 
        "content": (
            "Bạn là WASU's bot - trợ lý AI kỹ thuật và học tập chuyên nghiệp, sắc bén. "
            "QUY TẮC PHẢN HỒI BẮT BUỘC:\n"
            "1. ĐI THẲNG VÀO TRỌNG TÂM: Trả lời ngắn gọn, súc tích, ĐÚNG TRỌNG TÂM câu hỏi của người dùng. "
            "TUYỆT ĐỐI KHÔNG dài dòng, không lan man, không lặp ý hoặc dùng các lời dẫn xã giao sáo rỗng.\n"
            "2. ĐỊNH DẠNG TOÁN HỌC (LATEX): Khi viết công thức toán học, lý thuyết hay ký hiệu vật lý, "
            "LUÔN LUÔN dùng $...$ cho công thức inline và $$...$$ cho khối phương trình độc lập. "
            "TUYỆT ĐỐI KHÔNG dùng dấu ngoặc kép bọc LaTeX như ((...)) hay hiển thị sai ký tự.\n"
            "3. TRÌNH BÀY: Dùng Markdown rõ ràng (tiêu đề, danh sách, mã nguồn) để dễ đọc."
        )
    }

    # Đưa system prompt lên đầu tiên nếu chưa có
    if not messages_history or messages_history[0].get("role") != "system":
        messages_history.insert(0, system_prompt)

    try:
        completion = client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=messages_history,
            temperature=0.3,          # Giảm xuống 0.3 để hạn chế lan man, bám sát vấn đề
            max_completion_tokens=1500, # Giới hạn độ dài tránh trả lời quá đà
            top_p=1,
            stream=False,
            stop=None
        )
        bot_reply = completion.choices[0].message.content
    except Exception as e:
        bot_reply = f"Lỗi Groq API: {str(e)}"
    
    return jsonify({'reply': bot_reply})


# ==========================================
# API XUẤT FILE TỪ NỘI DUNG CHAT
# ==========================================
@app.route('/api/export/word', methods=['POST'])
def export_word():
    try:
        data = request.get_json() or {}
        content = data.get('content', 'Không có nội dung')
        
        doc = Document()
        doc.add_heading("WASU's Bot - Báo Cáo", level=1)
        
        for line in content.split('\n'):
            if line.strip():
                doc.add_paragraph(line)
        
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        
        return send_file(
            file_stream,
            as_attachment=True,
            download_name='wasu_bao_cao.docx',
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/export/excel', methods=['POST'])
def export_excel():
    try:
        data = request.get_json() or {}
        content = data.get('content', '')
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DuLieuBot"
        
        ws.append(["Nội dung kết quả từ WASU's Bot"])
        for line in content.split('\n'):
            if line.strip():
                ws.append([line.strip()])
        
        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)
        
        return send_file(
            file_stream,
            as_attachment=True,
            download_name='wasu_du_lieu.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except Exception as e:
        return jsonify({'error': str(e)}), 500


if __name__ == '__main__':
    app.run(host='127.0.0.1', port=5000, debug=True)